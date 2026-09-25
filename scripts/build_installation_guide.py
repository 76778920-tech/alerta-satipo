"""Genera Word y PDF de instalación a partir de Markdown versionado."""
from pathlib import Path
from html import escape
import textwrap
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.platypus import SimpleDocTemplate, Paragraph, Image, PageBreak, Preformatted, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'docs';ASSETS=OUT/'instalacion_figuras';ASSETS.mkdir(exist_ok=True)
fig,ax=plt.subplots(figsize=(12,6));fig.patch.set_facecolor('#F3F7F5');ax.axis('off');ax.set(xlim=(0,12),ylim=(0,6))
ax.text(.3,5.6,'Elegir la ruta de instalación',size=22,weight='bold',color='#205C45')
def box(x,y,w,h,title,body):
 ax.add_patch(FancyBboxPatch((x,y),w,h,boxstyle='round,pad=.04',edgecolor='#93B5A1',facecolor='white'))
 ax.text(x+.2,y+h-.25,title,size=13,weight='bold',va='top',color='#205C45')
 ax.text(x+.2,y+h-.7,body,size=11,va='top',linespacing=1.5,color='#18362E')
box(.3,2.7,5.3,2.1,'A / Proyecto compartido existente','Obtener acceso autorizado\nConfigurar URL + clave pública\nNo ejecutar migraciones ni importar datos')
box(6.3,2.7,5.3,2.1,'B / Proyecto independiente vacío','Configurar su propia URL y claves\nSQL 001 → SQL 002 → seed.sql → SQL 003\nAprovisionar acceso administrativo autorizado')
box(2.9,.3,6.2,1.3,'Resultado común','npm start → login administrativo → verificar 300 lecturas')
for x in [3,9]:ax.annotate('',xy=(6,1.65),xytext=(x,2.65),arrowprops={'arrowstyle':'->','color':'#205C45','lw':2})
fig.savefig(ASSETS/'rutas_instalacion.png',dpi=200,bbox_inches='tight');plt.close(fig)

styles=getSampleStyleSheet();styles.add(ParagraphStyle('Body2',fontName='Helvetica',fontSize=10,leading=14,spaceAfter=8))
styles.add(ParagraphStyle('Code2',fontName='Courier',fontSize=8,leading=10,spaceAfter=10,backColor=colors.HexColor('#F0F5F2')))
for name in ['Title','Heading1','Heading2']:styles[name].textColor=colors.HexColor('#205C45')
doc=Document();sec=doc.sections[0];sec.top_margin=sec.bottom_margin=Inches(.65)
normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10)
for name in ['Heading 1','Heading 2']:doc.styles[name].font.color.rgb=RGBColor.from_string('205C45')
sec.header.paragraphs[0].text='ALERTA SATIPO / GUÍA DE INSTALACIÓN'
foot=sec.footer.paragraphs[0];foot.alignment=2;foot.add_run('Instalación | ')
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');foot._p.append(f)
story=[]
def para(t,style='Body2'):story.append(Paragraph(escape(t),styles[style]))
doc.add_heading('Alerta Satipo',0);doc.add_heading('Guía de instalación y configuración',1)
doc.add_paragraph('Windows · PowerShell · Supabase\nVersión 1.0 / 25 de septiembre de 2026')
doc.add_picture(str(ASSETS/'rutas_instalacion.png'),width=Inches(6))
doc.add_paragraph('Procedimiento para ejecutar el proyecto existente o preparar un entorno independiente. Incluye verificaciones, resolución de errores y publicación opcional.')
doc.add_page_break()
para('ALERTA SATIPO','Title');para('Guía de instalación y configuración','Heading1')
para('Windows · PowerShell · Supabase | Versión 1.0 | 25 de septiembre de 2026')
story.append(Spacer(1,20));story.append(Image(str(ASSETS/'rutas_instalacion.png'),width=480,height=240))
para('Figura 1. Rutas de instalación. Diagrama explicativo, no captura de ejecución.')
para('Procedimiento para ejecutar el proyecto existente o preparar un entorno independiente. Incluye verificaciones, resolución de errores y publicación opcional.')
story.append(PageBreak())
lines=(OUT/'GUIA_INSTALACION.md').read_text(encoding='utf-8').splitlines()
para('Contenido','Heading1');doc.add_heading('Contenido',1)
for line in lines:
 if line.startswith('## '):para(line[3:]);doc.add_paragraph(line[3:])
story.append(PageBreak());doc.add_page_break()
code=False;buf=[];first=True
for line in lines:
 if line.startswith('# ') or line.startswith('Versión 1.0'):continue
 if line.startswith('```'):
  if code:
   for row in buf:
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);r=p.add_run(row);r.font.name='Consolas';r.font.size=Pt(8)
   wrapped='\n'.join('\n'.join(textwrap.wrap(row,96,replace_whitespace=False,drop_whitespace=False)) or ' ' for row in buf)
   story.append(Preformatted(wrapped,styles['Code2']));buf=[]
  code=not code;continue
 if code:buf.append(line);continue
 if not line:continue
 if line.startswith('## '):
  if not first:doc.add_page_break();story.append(PageBreak())
  first=False;doc.add_heading(line[3:],1);para(line[3:],'Heading1')
 elif line.startswith('### '):
  if line.startswith('### Comprobar'):
   doc.add_page_break();story.append(PageBreak())
  doc.add_heading(line[4:],2);para(line[4:],'Heading2')
 elif line.startswith('- '):doc.add_paragraph(line[2:],style='List Bullet');para('• '+line[2:])
 else:doc.add_paragraph(line);para(line)
doc.save(OUT/'GUIA_INSTALACION_ALERTA_SATIPO.docx')
def footer(c,d):
 c.setFont('Helvetica',8);c.setFillColor(colors.HexColor('#526B60'));c.drawString(45,817,'ALERTA SATIPO / GUÍA DE INSTALACIÓN');c.drawRightString(550,25,str(d.page))
SimpleDocTemplate(str(OUT/'GUIA_INSTALACION_ALERTA_SATIPO.pdf'),pagesize=(595,842),leftMargin=45,rightMargin=45,topMargin=45,bottomMargin=45).build(story,onFirstPage=footer,onLaterPages=footer)
print('Guía Word/PDF generada desde docs/GUIA_INSTALACION.md')
