"""Láminas y memoria de diseño propuesto; no modifica el backend desplegado."""
from pathlib import Path
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, Polygon
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from docx import Document
from docx.shared import Inches
from html import escape

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'docs'/'arquitectura_hexagonal';OUT.mkdir(exist_ok=True)
INK='#172C40'; BLUE='#245D87'; GREEN='#206653'; GREY='#526779'
def base(title,sub):
 f,a=plt.subplots(figsize=(18,10));a.set(xlim=(0,18),ylim=(0,10));a.axis('off');f.patch.set_facecolor('#F6F8FB')
 a.text(.4,9.5,title,size=24,weight='bold',color=INK)
 a.text(.4,9.05,sub,size=12,color=GREY)
 return f,a
def box(a,x,y,w,h,title,body,color=BLUE):
 a.add_patch(FancyBboxPatch((x,y),w,h,boxstyle='round,pad=.04,rounding_size=.12',edgecolor=color,facecolor='white',linewidth=1.3))
 a.text(x+.15,y+h-.23,title,va='top',size=11,weight='bold',color=color)
 a.text(x+.15,y+h-.7,body,va='top',size=10.5,linespacing=1.5,color=INK)
def arrow(a,start,end,label='',dashed=False):
 a.annotate('',xy=end,xytext=start,arrowprops=dict(arrowstyle='-|>',color=BLUE,lw=1.6,linestyle='--' if dashed else '-'))
 if label:a.text((start[0]+end[0])/2,(start[1]+end[1])/2+.12,label,size=9,ha='center',color=GREY,bbox=dict(facecolor='#F6F8FB',edgecolor='none',pad=1))
def save(f,name):
 for ext in ['png','svg','pdf']:f.savefig(OUT/f'{name}.{ext}',dpi=210,bbox_inches='tight')
 plt.close(f)

f,a=base('ALERTA SATIPO / Arquitectura hexagonal propuesta','VISTA ESTÁTICA · Dependencias de código y contratos · Diseño objetivo, no implementación actual')
a.add_patch(Polygon([(4.4,4.3),(5.5,8.25),(12.4,8.25),(13.6,4.3),(12.4,1.8),(5.5,1.8)],closed=True,facecolor='#E9F2EF',edgecolor=GREEN,lw=2))
a.text(9,7.92,'NÚCLEO INDEPENDIENTE DE FRAMEWORKS',ha='center',size=12,weight='bold',color=GREEN)
box(a,.4,5.5,3.5,2.3,'Adaptador de entrada / HTTP','Node.js + Express (propuesto)\nValida formato y mapea errores\nInvoca el puerto de entrada\nNo decide reglas de dominio')
box(a,.4,2.6,3.5,1.8,'Adaptador de pruebas','Pruebas de casos de uso\nEntrada directa, sin HTTP\nRepositorios en memoria')
box(a,5.5,5.6,3,1.7,'Puerto de entrada','ActualizarMantenimiento\nConsultarLecturas\nRevisarCaso',GREEN)
box(a,9.5,5.6,3,1.7,'Aplicación','Implementa casos de uso\nAutoriza al actor\nCoordina dominio y puertos',GREEN)
box(a,5.5,2.45,3,1.9,'Dominio','Tarea · Caso · Lectura\nEstados permitidos\nIdentidad y reglas\nSin SDK, HTTP ni SQL',GREEN)
box(a,9.5,2.45,3,1.9,'Puertos de salida','RepositorioMantenimiento\nConsultaLecturas\nIdentidadYAutorización\nContratos del núcleo',GREEN)
box(a,14.1,5.5,3.5,2.3,'Adaptador / persistencia','Implementa repositorios\nSDK supabase-js → API\nMapeo filas ↔ objetos\nPropaga JWT de usuario')
box(a,14.1,2.6,3.5,1.8,'Adaptador / identidad','Valida sesión en Supabase\nConsulta rol autorizado\nNo confía en rol del navegador')
arrow(a,(3.9,6.4),(5.5,6.4),'usa')
arrow(a,(3.9,3.6),(5.5,5.6),'invoca')
arrow(a,(9.5,6.5),(8.5,6.5),'implementa',True)
arrow(a,(10.5,5.6),(7.6,4.35),'usa')
arrow(a,(11.5,5.6),(11.5,4.35),'usa')
arrow(a,(14.1,5.6),(12.5,4.1),'implementa',True)
arrow(a,(14.1,3.2),(12.5,3.2),'implementa',True)
a.text(.4,1.15,'REGLA DE DEPENDENCIA',size=11,weight='bold',color=GREEN)
a.text(.4,.72,'Las flechas señalan el elemento del que se depende. Los adaptadores dependen de contratos del núcleo; el núcleo no importa adaptadores.',size=11,color=INK)
a.text(.4,.28,'Línea continua: uso   ·   Línea discontinua: implementación de contrato   ·   Composition root: ensambla e inyecta adaptadores al arrancar.',size=10.5,color=GREY)
save(f,'01_dependencias_hexagonales')

f,a=base('ALERTA SATIPO / Actualizar mantenimiento','VISTA DINÁMICA · Flujo propuesto de una solicitud · Las flechas representan mensajes, no dependencias de código')
xs=[1.6,5.1,8.7,12.2,16.2]
names=['Panel web','HTTP + caso de uso','Dominio','Adaptador Supabase','PostgreSQL / RLS']
for x,n in zip(xs,names):
 box(a,x-1.45,7.6,2.9,.8,n,'');a.plot([x,x],[1.45,7.55],color='#B4C5D0',ls='--',lw=1)
events=[(0,1,7.15,'1. PATCH + JWT + estado esperado'),(1,3,6.5,'2. Verificar identidad y autorización'),(3,1,5.85,'3. Actor verificado / rechazo'),(1,2,5.2,'4. Validar nuevo estado'),(2,1,4.55,'5. Estado válido / error de dominio'),(1,3,3.9,'6. guardarSiCoincide(id, esperado, nuevo)'),(3,4,3.25,'7. UPDATE condicional + RLS'),(4,3,2.65,'8. Resultado SQL'),(3,1,2.05,'9. Resultado del repositorio'),(1,0,1.5,'10. Respuesta HTTP')]
for s,t,y,label in events:arrow(a,(xs[s],y),(xs[t],y),label)
a.text(.4,1.0,'RESPUESTAS PROPUESTAS',size=12,weight='bold',color=GREEN)
a.text(.4,.6,'200: cambio persistido   |   401: sesión inválida   |   403: rol insuficiente   |   404: tarea inexistente',size=12,color=INK)
a.text(.4,.2,'409: estado esperado no coincide   |   422: estado no permitido   |   503: dependencia no disponible',size=12,color=INK)
save(f,'02_secuencia_mantenimiento')

sections=[
('1. Propósito, alcance y estado del diseño',
 'Esta memoria propone una evolución del backend de Alerta Satipo hacia puertos y adaptadores. La versión vigente es un frontend estático que consulta Supabase directamente mediante JavaScript. No existen todavía un servidor Express, contratos de repositorios ni casos de uso independientes. Las láminas describen el objetivo, no certifican su implementación.\nEl alcance propuesto es el panel administrativo: consulta de las 300 lecturas, revisión de casos y actualización del mantenimiento demostrativo. Se conservan Supabase Auth y PostgreSQL. Los nodos siguen siendo seis lotes virtuales de 50 lecturas; las etiquetas positivas no se convierten en incendios confirmados.'),
('2. Decisión arquitectónica y responsabilidades',
 'Decisión ADR-001: introducir un servicio de aplicación en Node.js y Express, con dependencias dirigidas hacia contratos propios. Motivación: extraer la coordinación de operaciones del navegador, centralizar autorización y probar reglas sin red. Coste: nuevo despliegue, configuración de CORS y TLS, latencia adicional, gestión de errores y mantenimiento operativo. Para un prototipo pequeño, Supabase directo sigue siendo una solución válida; la evolución se justifica por requisitos académicos y de mantenibilidad, no por una supuesta incapacidad de Supabase.\nDominio: objetos y reglas expresados en JavaScript puro. Aplicación: casos de uso, autorización y coordinación. Puertos de entrada: operaciones que ofrece el núcleo. Puertos de salida: capacidades externas que necesita. Adaptadores: HTTP, persistencia e identidad Supabase. El punto de composición crea las implementaciones e inyecta las dependencias.\nLa inversión de dependencias ocurre porque el repositorio Supabase implementa un contrato definido por el núcleo. Que el caso de uso invoque al repositorio durante la ejecución no implica que su código dependa de la clase Supabase. Las dos láminas separan deliberadamente estructura y ejecución.'),
('3. Contratos propuestos y reglas verificables',
 'Puerto de entrada: actualizarMantenimiento({actor, nodeId, expectedState, nextState}) devuelve una tarea actualizada o un error tipado. actor debe provenir de una identidad verificada, no de un campo enviado libremente por el navegador.\nPuerto de persistencia: obtener(nodeId) devuelve Tarea o ausencia; guardarSiCoincide(nodeId, expectedState, nextState) devuelve actualizado o conflicto, y distingue errores de infraestructura. Puerto de identidad: verificarSesion(token) y esAdministrador(userId). El adaptador HTTP no conoce tablas; el caso de uso no conoce códigos HTTP; el adaptador de persistencia traduce filas y errores del proveedor.\nInvariantes existentes: state pertenece a Pendiente, En progreso o Completada; node_id identifica un nodo existente; una tarea se asocia a un solo nodo. Las restricciones actuales permiten cambiar entre cualquiera de esos estados. No se debe afirmar que hay una máquina de transiciones estricta o prohibición de reabrir tareas: esas serían reglas nuevas que requieren definición.\nConcurrencia: el UPDATE filtra por node_id y state esperado; evita sobrescribir un cambio incompatible. Si devuelve cero filas, se consulta existencia con el mismo contexto autorizado para distinguir ausencia de conflicto. Comparar solo el estado no detecta un cambio A→B→A; para garantizar detección completa se propone una columna version e incremento atómico mediante operación SQL/RPC. Esa versión no está implementada.\nAtomicidad: actualizar una tarea es una operación SQL. Si se agrega una auditoría obligatoria de mantenimiento, el cambio y el evento deben confirmarse en la misma transacción, mediante función o disparador. Hoy solo los incidentes tienen auditoría; no se afirma que las tablas demo dispongan de ella.'),
('4. Seguridad y persistencia',
 'El adaptador de identidad debe validar la sesión en Supabase y resolver la autorización administrativa en el servidor. No basta decodificar un JWT ni leer sessionStorage. La respuesta nunca debe incluir claves privadas o tokens en registros.\nPara persistencia se propone usar clave pública y JWT del usuario en un cliente acotado a la solicitud. Así PostgreSQL conserva RLS como segunda barrera. No utilizar una sesión mutable global compartida entre usuarios. Usar una clave administrativa en las solicitudes ordinarias podría eludir RLS y exigiría controles equivalentes explícitos; no es la opción propuesta.\nEl traslado al servidor no elimina automáticamente las rutas REST directas de Supabase. Mientras existan los privilegios actuales, un administrador puede seguir modificando state por esa API. Si una regla debe ser imposible de eludir, también debe imponerse en PostgreSQL o restringirse la escritura a una operación controlada. CORS no sustituye esos permisos.\nLecturas históricas: claves dataset_id y source_row preservan la trazabilidad. El puerto de consulta debe incluir filtros y paginación sin devolver objetos del SDK al dominio. La importación de CSV sigue siendo una tarea administrativa separada; no se descarga GitHub en cada consulta del panel.'),
('5. Migración, trazabilidad y aceptación',
 'Paso 1: crear backend/src/domain, application/ports, application/use-cases, adapters/in/http, adapters/out/supabase y bootstrap. Los directorios solo organizan; la independencia se demuestra con importaciones y pruebas.\nPaso 2: extraer la actualización de mantenimiento de frontend/web/js/operations.js a un caso de uso. Definir los contratos antes de implementar el adaptador. Mantener las reglas SQL de la migración 003.\nPaso 3: implementar adaptadores HTTP, identidad y persistencia; trasladar la invocación del panel al nuevo endpoint. Configurar alojamiento del servicio, HTTPS, origen permitido y variables privadas. Firebase Hosting actual publica archivos, no el nuevo proceso Express.\nPaso 4: migrar consultas de lecturas desde frontend/shared/js/dataset.js y casos desde operations.js; no reescribir todo en una sola entrega. Mantener verificaciones de las 300 filas y los vínculos existentes.\nAceptación: pruebas unitarias del dominio sin SDK; casos de uso con repositorio en memoria; pruebas de contrato que ejecuten la misma suite sobre memoria y Supabase de prueba; pruebas HTTP de 401/403/404/409/422/503; integración que demuestre RLS; dos actualizaciones concurrentes con un mismo estado esperado y destino distinto; prueba de regresión que conserve los 300 registros.\nAñadir comprobación automatizada de importaciones: domain no importa Express, Supabase ni módulos de adapters; application no importa implementaciones de infraestructura. La prueba debe fallar si se introduce una dependencia prohibida. No se han ejecutado estas pruebas propuestas porque el servicio todavía no existe.'),
('6. Guion de sustentación y evidencia',
 'Explicación: “Separamos reglas, coordinación e infraestructura. El núcleo define los contratos; HTTP inicia los casos de uso y Supabase implementa persistencia e identidad. La dirección de las dependencias apunta al núcleo, aunque el flujo de ejecución salga hacia la base de datos. Conservamos RLS y restricciones SQL como garantías adicionales. Esta es una propuesta de evolución verificable, no una descripción ficticia del despliegue actual”.\nSi preguntan por qué usar hexagonal: permite sustituir adaptadores y probar casos de uso sin red, a cambio de mayor complejidad. Si preguntan si cambiar de base de datos es gratis: no; se reemplaza el adaptador y se migran esquema, políticas y funciones, aunque los casos de uso deberían permanecer estables. Si preguntan qué prueba la arquitectura: los contratos, la composición, las importaciones y las pruebas, no la forma del dibujo.\nFuentes internas: frontend/shared/js/backend.js (identidad y guard); frontend/web/js/operations.js (estados y actualización condicional); frontend/shared/js/dataset.js (lecturas); supabase/migrations/202609230003_demo_operations.sql (relaciones y RLS); tests/database.test.mjs y tests/panel_operations.py (verificación actual). Estos archivos respaldan el punto de partida; no demuestran la implementación del diseño futuro.')]

md='# Alerta Satipo — Memoria de arquitectura hexagonal propuesta\n\nFecha: 25 de septiembre de 2026 · Estado: diseño pendiente de implementación.\n\n'
for title,body in sections:md+='## '+title+'\n\n'+body.replace('\n','\n\n')+'\n\n'
(OUT/'MEMORIA_ARQUITECTURA.md').write_text(md,encoding='utf-8')
styles=getSampleStyleSheet();styles.add(ParagraphStyle('Body2',fontName='Helvetica',fontSize=10.5,leading=15,spaceAfter=10))
styles['Heading1'].textColor=colors.HexColor(GREEN)
story=[];doc=Document()
doc.add_heading('Alerta Satipo',0);doc.add_heading('Arquitectura hexagonal propuesta',1)
doc.add_paragraph('Memoria de diseño · 25 de septiembre de 2026\nEstado: pendiente de implementación')
for title,body in sections:
 if story:story.append(PageBreak())
 story.append(Paragraph(escape(title),styles['Heading1']))
 doc.add_heading(title,1)
 for p in body.split('\n'):
  story.append(Paragraph(escape(p),styles['Body2']));doc.add_paragraph(p)
 if title.startswith('1.'):
  for name in ['01_dependencias_hexagonales','02_secuencia_mantenimiento']:
   doc.add_page_break();doc.add_picture(str(OUT/f'{name}.png'),width=Inches(6.2))
  doc.add_page_break()
doc.save(OUT/'MEMORIA_ARQUITECTURA.docx')
def footer(c,d):
 c.setFont('Helvetica',8);c.setFillColor(colors.HexColor(GREY));c.drawString(45,25,'ALERTA SATIPO / PROPUESTA DE DISEÑO');c.drawRightString(550,25,str(d.page))
SimpleDocTemplate(str(OUT/'MEMORIA_ARQUITECTURA.pdf'),pagesize=(595,842),leftMargin=45,rightMargin=45,topMargin=40,bottomMargin=45).build(story,onFirstPage=footer,onLaterPages=footer)
print('Generadas dos laminas PNG/SVG/PDF y memoria Word/PDF/Markdown.')
