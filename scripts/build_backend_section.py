"""Exporta el apartado 2 vigente a Word y PDF, sin incluir otros capítulos."""
from pathlib import Path
from html import escape
import textwrap
from docx import Document
from docx.shared import Pt, Inches
from reportlab.platypus import SimpleDocTemplate, Paragraph, Preformatted, KeepTogether
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
root=Path(__file__).resolve().parents[1]/'docs'
doc=Document();doc.sections[0].top_margin=doc.sections[0].bottom_margin=Inches(.7)
doc.styles['Normal'].font.name='Calibri';doc.styles['Normal'].font.size=Pt(11)
styles=getSampleStyleSheet()
styles.add(ParagraphStyle('Body2',fontSize=10.5,leading=15,spaceAfter=9))
styles.add(ParagraphStyle('Code2',fontName='Courier',fontSize=8,leading=11,spaceAfter=10,backColor=colors.HexColor('#F1F5F3')))
story=[];code=False;buf=[]
for line in (root/'PARTE_2_BACKEND_INSTALACION.md').read_text(encoding='utf-8').splitlines():
    if line.startswith('```'):
        if code:
            for i,row in enumerate(buf):
                p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.keep_with_next=i<len(buf)-1
                r=p.add_run(row);r.font.name='Consolas';r.font.size=Pt(8)
            wrapped='\n'.join('\n'.join(textwrap.wrap(row,94,replace_whitespace=False,drop_whitespace=False)) or ' ' for row in buf)
            story.append(KeepTogether([Preformatted(wrapped,styles['Code2'])]));buf=[]
        code=not code;continue
    if code:buf.append(line);continue
    if not line:continue
    if line.startswith('# '):doc.add_heading(line[2:],1);story.append(Paragraph(escape(line[2:]),styles['Heading1']))
    elif line.startswith('## '):doc.add_heading(line[3:],2);story.append(Paragraph(escape(line[3:]),styles['Heading2']))
    elif line.startswith('- '):doc.add_paragraph(line[2:],style='List Bullet');story.append(Paragraph('• '+escape(line[2:]),styles['Body2']))
    else:doc.add_paragraph(line);story.append(Paragraph(escape(line),styles['Body2']))
doc.save(root/'PARTE_2_BACKEND_HEXAGONAL.docx')
def footer(c,d):
    c.setFont('Helvetica',8);c.drawString(45,815,'ALERTA SATIPO / BACKEND HEXAGONAL IMPLEMENTADO');c.drawRightString(550,25,str(d.page))
SimpleDocTemplate(str(root/'PARTE_2_BACKEND_HEXAGONAL.pdf'),pagesize=(595,842),leftMargin=45,rightMargin=45,topMargin=45,bottomMargin=45).build(story,onFirstPage=footer,onLaterPages=footer)
print('Apartado 2 hexagonal exportado.')
